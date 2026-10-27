"""The website's Excel export (templates/xlsx.js): the workbook it writes opens, holds the data, the chart drawn from
those cells, and the About sheet's links. The writer is run in Node; the page code that calls it needs a browser."""

from __future__ import annotations

import json
import shutil
import subprocess
import zipfile
from importlib import resources

import openpyxl
import pytest

NODE = shutil.which("node")
needs_node = pytest.mark.skipif(NODE is None, reason="Node.js not installed")

HARNESS = r"""
const fs = require("fs"), vm = require("vm");
const [src, specFile, out] = process.argv.slice(2);
vm.runInThisContext(fs.readFileSync(src, "utf8"));
const spec = JSON.parse(fs.readFileSync(specFile, "utf8"));
const names = xlSheetNames([...spec.tables.map((t) => t.name), "About"]);
const sheets = [...spec.tables.map((t, i) => xlTableSheet({ ...t, name: names[i] })), xlAboutSheet(spec.about)];
fs.writeFileSync(out, xlsxWorkbook({ title: spec.about.title, sheets }, new Date(2026, 9, 27, 9, 30)));
console.log(JSON.stringify(xlSheetNames(spec.names)));
"""

YEARS = ["2014-15", "2015-16", "2016-17", "2017-18"]
SPEC = {
    "tables": [
        {
            "name": "Chart",
            "title": "Return on equity by group",
            "unit": "Return on equity after tax, %.",
            "lines": ["Settings: Period: 2014-15 to 2017-18."],
            "head": ["Financial year", "Banks & insurers <test>", "Low barriers"],
            "rows": [[y, 10 + i, None if i == 2 else 8.5 - i] for i, y in enumerate(YEARS)],
            "fmts": ["text", "0.0", "0.0"],
            "chart": {
                "type": "line",
                "cat": 0,
                "yMin": 0,
                "yMax": 15,
                "series": [
                    {"col": 1, "color": "#2a78d6", "width": 1.5},
                    {"col": 2, "color": "#6b6a65", "dash": "dash"},
                ],
            },
        },
        {
            "name": "Table",
            "title": "Growth by state",
            "head": ["Industry", "NSW", "VIC"],
            "rows": [["Mining", 3.2, -1.5], ["Retail", -0.4, 2.0]],
            "fmts": ["text", "+0.0;-0.0;0.0", "+0.0;-0.0;0.0"],
            "colorScale": {"kind": "div", "scale": 3, "neg": "#e8a19c", "mid": "#fcfcfb", "pos": "#9fcac3"},
        },
    ],
    "about": {
        "title": "Return on equity by group",
        "unit": "Return on equity after tax, %.",
        "page": "https://example.github.io/site/conditions.html",
        "exported": "27 October 2026",
        "built": "2026-10-27",
        "settings": [["Period", "2014-15 to 2017-18"]],
        "sources": [
            {
                "label": "ATO, Taxation Statistics",
                "url": "https://data.gov.au/data/dataset/taxation-statistics-2022-23",
                "release": "2022-23",
                "licence": "CC BY 2.5 AU",
            }
        ],
        "notes": [["About these figures", "Companies only."]],
        "terms": [["Equity", "Assets less liabilities."]],
        "links": [["Methods note", "https://github.com/example/docs/methods.md"]],
        "caveat": "Descriptive only.",
    },
    "names": ["Has concentration risen?", "Has concentration risen?", "A [very] long name: with/odd*characters, more"],
}


@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    if NODE is None:
        pytest.skip("Node.js not installed")
    d = tmp_path_factory.mktemp("xlsx")
    src = resources.files("alicia.menu").joinpath("templates/xlsx.js")
    (d / "spec.json").write_text(json.dumps(SPEC), encoding="utf-8")
    (d / "run.js").write_text(HARNESS, encoding="utf-8")
    out = d / "export.xlsx"
    done = subprocess.run(
        [NODE, str(d / "run.js"), str(src), str(d / "spec.json"), str(out)],
        capture_output=True,
        text=True,
        check=True,
        timeout=60,
    )
    return out, json.loads(done.stdout)


@needs_node
def test_export_is_a_valid_workbook_with_the_data_as_numbers(workbook):
    path, _ = workbook
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        assert "[Content_Types].xml" in z.namelist() and "xl/charts/chart1.xml" in z.namelist()
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Chart", "Table", "About"]
    ws = wb["Chart"]
    assert ws["A1"].value == "Return on equity by group"
    head = next(r for r in ws.iter_rows(values_only=True) if r[0] == "Financial year")
    assert head == ("Financial year", "Banks & insurers <test>", "Low barriers")
    rows = [r for r in ws.iter_rows(values_only=True) if r[0] in YEARS]
    assert [r[1] for r in rows] == [10, 11, 12, 13]
    assert rows[2][2] is None  # a missing value is a blank cell, so Excel leaves a gap
    cell = next(c for c in ws["B"] if c.value == 10)
    assert cell.number_format == "0.0"


@needs_node
def test_export_chart_is_drawn_from_the_sheet_cells(workbook):
    path, _ = workbook
    ws = openpyxl.load_workbook(path)["Chart"]
    assert len(ws._charts) == 1
    header_row = next(c.row for c in ws["A"] if c.value == "Financial year")
    first, last = header_row + 1, header_row + len(YEARS)
    series = ws._charts[0].series
    assert [s.val.numRef.f for s in series] == [f"'Chart'!$B${first}:$B${last}", f"'Chart'!$C${first}:$C${last}"]
    assert series[0].cat.strRef.f == f"'Chart'!$A${first}:$A${last}"
    assert series[0].tx.strRef.f == f"'Chart'!$B${header_row}"
    xml = zipfile.ZipFile(path).read("xl/charts/chart1.xml").decode()
    assert '<a:prstDash val="dash"/>' in xml and '<c:dispBlanksAs val="gap"/>' in xml
    assert '<c:max val="15"/>' in xml and '<c:min val="0"/>' in xml


@needs_node
def test_export_colours_a_table_like_the_page_heatmap(workbook):
    path, _ = workbook
    ws = openpyxl.load_workbook(path)["Table"]
    rules = [r for cf in ws.conditional_formatting for r in cf.rules]
    assert len(rules) == 1 and rules[0].type == "colorScale"
    assert [float(v.val) for v in rules[0].colorScale.cfvo] == [-3, 0, 3]


@needs_node
def test_export_about_sheet_links_each_source_with_its_licence(workbook):
    path, _ = workbook
    ws = openpyxl.load_workbook(path)["About"]
    cells = {c.value: c for row in ws.iter_rows() for c in row if c.value is not None}
    head = cells["Dataset"].row
    assert [ws.cell(head, c).value for c in range(1, 5)] == ["Dataset", "Release used", "Licence", "Link"]
    assert ws.cell(head + 1, 3).value == "CC BY 2.5 AU"
    link = ws.cell(head + 1, 4)
    assert link.hyperlink.target == "https://data.gov.au/data/dataset/taxation-statistics-2022-23"
    assert cells["https://github.com/example/docs/methods.md"].hyperlink is not None
    for text in ("Settings on the page when exported", "Notes", "Words used", "Method and limits", "Descriptive only."):
        assert text in cells


@needs_node
def test_sheet_names_are_valid_and_unique(workbook):
    _, names = workbook
    assert names[0] == "Has concentration risen" and names[1] != names[0]
    assert all(len(n) <= 31 and not set(n) & set(":\\/?*[]") for n in names)
