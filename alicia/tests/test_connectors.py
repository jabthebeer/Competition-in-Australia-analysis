"""Connector tests on small synthetic inputs (ILLUSTRATIVE values, not real data)."""

from __future__ import annotations

import datetime as dt

import pandas as pd
import pytest

from alicia import anzsic
from alicia.pipeline import schema


def test_name_rules_use_real_codes():
    from alicia.pipeline import names

    for code, _ in names.RULES + names.OVERRIDES:
        assert code == "" or anzsic.name(code) is not None, code


def test_classify_overrides_and_leftmost_match():
    from alicia.pipeline import names

    assert names.classify("Harbour Pharmacy") == "4271"
    assert names.classify("Smith Plumbing & Electrical") == "3231"  # leftmost rule wins
    assert names.classify("Plumbing Supplies Co") == "F"  # supplier override wins anywhere
    assert names.classify("Master Builders Association") == ""
    assert names.classify("Arcadia Landscape Architects") == "6921"
    assert names.classify("Sushi Ten") == "451"
    assert names.classify("Hardware Lane Bistro") == "4511"
    assert names.classify("SMITHPLUMBING") == ""  # whole words only
    assert names.classify(None, "", "Bay Taxis") == "4623"
    assert [names.level(c) for c in ("E", "451", "4271", "")] == ["division", "group", "class", ""]


def test_classifier_validation_counts():
    from alicia.pipeline import names

    r = names.validate(["Bay Pharmacy", "Joe's Cafe", "Nothing Here", "Sushi Go"], ["4271", "4512", "9999", "4512"])
    assert (r["classified"], r["coverage"]) == (3, 0.75)
    assert r["by_code"]["4511"]["precision"] == 0.0 and r["by_code"]["451"]["precision"] == 1.0


def test_abr_line_parse():
    from alicia.pipeline import abr_abn

    line = (
        '<ABR recordLastUpdatedDate="20240102" replaced="N"><ABN status="ACT" ABNStatusFromDate="20200301">'
        "11111111111</ABN><EntityType><EntityTypeInd>PRV</EntityTypeInd><EntityTypeText>Australian Private "
        'Company</EntityTypeText></EntityType><MainEntity><NonIndividualName type="MN"><NonIndividualNameText>'
        "EXAMPLE HOLDINGS PTY LTD</NonIndividualNameText></NonIndividualName><BusinessAddress><AddressDetails>"
        "<State>WA</State><Postcode>6000</Postcode></AddressDetails></BusinessAddress></MainEntity>"
        '<GST status="ACT" GSTStatusFromDate="20200401" /><OtherEntity><NonIndividualName type="BN">'
        "<NonIndividualNameText>EXAMPLE BAY PHARMACY</NonIndividualNameText></NonIndividualName></OtherEntity></ABR>"
    )
    r = abr_abn.parse_line(line)
    assert (r["abn"], r["status"], r["status_from"], r["group"]) == (11111111111, "ACT", 20200301, "company")
    assert (r["state"], r["postcode"], r["gst"], r["gst_from"]) == ("WA", "6000", "ACT", 20200401)
    assert r["industry"] == ""  # "HOLDINGS" in the main name leaves it unclassified
    assert abr_abn.parse_line("<ABR nothing>") is None
    assert abr_abn.entity_group("IND", "Individual/Sole Trader") == "individual"
    assert abr_abn.entity_group("SMF", "ATO Regulated Self-Managed Superannuation Fund") == "super fund"


def _toy(**over):
    base = dict(
        source="toy",
        measure="m",
        unit="count",
        adjustment="original",
        industry_scheme="none",
        industry_level="total",
        industry_code="",
        geo_level="POA",
        geo_code="6000",
        breakdown="",
        period_type="month",
        period=pd.Timestamp("2020-01-01"),
        value=10.0,
        vintage="t",
    )
    return {**base, **over}


def test_geo_apportion(monkeypatch):
    from alicia.pipeline import geo

    conc = pd.DataFrame(
        {
            "postcode": ["6000", "6000", "6001"],
            "level": "SA2",
            "code": ["501", "502", "502"],
            "name": "",
            "ratio": [0.75, 0.25, 1.0],
            "quality": "",
        }
    )
    monkeypatch.setattr(geo, "concordance", lambda: conc)
    panel = schema.make(pd.DataFrame([_toy(), _toy(geo_code="6001", value=4.0)]))
    out = geo.apportion(panel, "SA2").set_index("geo_code")["value"]
    assert out.to_dict() == {"501": 7.5, "502": 6.5}
    with pytest.raises(ValueError):
        geo.apportion(panel, "POA")


def test_asic_panel_and_complete_months():
    from alicia.pipeline import asic

    d = pd.DataFrame({"state": ["WA", "WA", ""], "postcode": ["6000", "6000", ""], "industry": ["4271", "", ""]})
    period = pd.Series(pd.Timestamp("2026-01-01"), index=d.index)
    parts = pd.concat(asic._panel(d, "x", period, None), ignore_index=True)
    got = {(r.geo_level, r.industry_code): r.value for r in parts.itertuples()}
    assert got == {("AUS", ""): 3, ("STE", ""): 2, ("POA", ""): 2, ("AUS", "4271"): 1, ("STE", "4271"): 1}
    dates = pd.Series(pd.to_datetime(["2012-05-31", "2012-06-01", "2026-09-30", "2026-10-02", None]))
    keep = asic._complete_months(dates, pd.Timestamp("2012-06-01"), pd.Timestamp("2026-10-14"))
    assert keep.tolist() == [False, True, True, False, False]


def test_lfs_tidy_handles_nfd_rows():
    from alicia.pipeline import abs_lfs

    month = dt.datetime(2026, 2, 1)
    rows = [
        (month, "Males", "Victoria", "412 Pharmaceutical and Other Store-Based Retailing", 1.0, 0.5, 40.0, 10.0),
        (month, "Females", "Victoria", "412 Pharmaceutical and Other Store-Based Retailing", 2.0, 1.0, 70.0, 20.0),
        (month, "Females", "Victoria", "410 Retail Trade nfd", 0.25, 0.0, 9.0, 0.0),
        (month, "Females", "Victoria", "G00 Retail Trade nfd", 0.5, 0.0, 15.0, 0.0),
    ]
    p = abs_lfs.tidy(pd.DataFrame(rows, columns=list(abs_lfs.COLUMNS)), "test")
    emp = p[(p["measure"] == "employed") & (p["geo_level"] == "STE")].set_index("industry_level")
    assert emp.loc["group", "industry_code"] == "412" and emp.loc["group", "value"] == 4.5
    assert emp.loc["subdivision", "value"] == 4.75 and emp.loc["division", "value"] == 5.25
    assert emp.loc["total", "value"] == 5.25
    assert set(p["period"]) == {pd.Timestamp("2026-01-01")} and set(p["period_type"]) == {"quarter"}
    hours = p[(p["measure"] == "hours_actual") & (p["industry_level"] == "group") & (p["geo_level"] == "AUS")]
    assert hours["value"].tolist() == [140.0]


def test_cabee_new_layouts():
    from alicia.outcomes import cabee

    assert cabee._fy_of_sheet("June 2021 a") == 2021 and cabee._fy_of_sheet("June 2021 b") is None
    title = "Businesses by Main State by Industry Class by Annualised Employment Size Ranges, June 2025 (a) (b)"
    assert cabee._fy_of_sheet("Table 1", title) == 2025
    assert cabee._fy_of_sheet("Table 2", title.replace("Annualised", "Experimental Point in Time")) is None
    heads = ["Operating at start of financial year", "Entries", "Exits", "Net Movement", "Operating at end"]
    bands = ["Non employing", "1-4 Employees", "5-19 Employees", "20-199 Employees", "200+ Employees", "Total"]
    raw = pd.DataFrame([[None] * 33, [None] * 33])
    for i, h in enumerate(heads):
        raw.iat[0, 3 + 6 * i] = h
        for j, b in enumerate(bands):
            raw.iat[1, 3 + 6 * i + j] = b
    cols = cabee._columns(raw, 0)
    assert cols["start|1-4"] == 4 and cols["end|total"] == 32 and "start|1-19" not in cols


def test_cabee_registered_sample():
    from alicia.outcomes import cabee

    df = pd.DataFrame(
        {
            "release": ["Jul 2016 to Jun 2020", "Jul 2016 to Jun 2020", "Jul 2017 to Jun 2021", "Jul 2021 to Jun 2025"],
            "fy": [2019, 2020, 2020, 2025],
        }
    )
    assert cabee.as_registered(df)["fy"].tolist() == [2019, 2020]


def test_clue_splits_map_to_anzsic_classes():
    from alicia.pipeline import clue

    assert all(anzsic.name(c) is not None and len(c) == 4 for c in clue.CLUE_SPLITS.values())
    assert clue.half("Example Cafe") == clue.half("EXAMPLE CAFE") in {"tuning", "holdout"}


def test_ivi_parser_on_synthetic_sheet():
    from alicia.pipeline import jsa_ivi

    raw = pd.DataFrame(
        [
            ["Internet Vacancy Index", None, None, None, None, None],
            ["Level", "ANZSCO_CODE", "ANZSCO_TITLE", "State", dt.datetime(2024, 1, 1), dt.datetime(2024, 2, 1)],
            [4, 3341, "Plumbers", "NSW", 120, 130],
            [4, 3341, "Plumbers", "Western Australia", 40, None],
            [0, 0, "Total", "AUST", 9000, 9100],
            [None, "Notes", None, None, None, None],
        ]
    )
    assert jsa_ivi.parse_sheet(raw.iloc[:, :4], "Trend", "f.xlsx") is None  # no month columns
    raw[6] = [None, dt.datetime(2024, 3, 1), 140, 50, 9200, None]
    p = jsa_ivi.parse_sheet(raw, "Trend", "f.xlsx")
    assert len(p) == 8 and set(p["adjustment"]) == {"trend"}
    wa = p[p["geo_code"] == "WA"].sort_values("period")
    assert wa["value"].tolist() == [40.0, 50.0] and set(wa["industry_level"]) == {"occupation"}
    assert set(p.loc[p["geo_level"] == "AUS", "industry_level"]) == {"total"}
    panel = schema.make(p, source="jsa_ivi", measure="online_vacancies", unit="count", period_type="month")
    assert len(panel) == 8


def test_regional_tidy():
    from alicia.pipeline import abs_regional

    raw = pd.DataFrame(
        {
            "DATAFLOW": "ABS:ABS_REGIONAL_ASGS2021(1.6.0)",
            "MEASURE": ["CABEE_34", "LEED_11", "CABEE_5", "CABEE_5"],
            "REGIONTYPE": ["SA2", "SA2", "STE", "STE"],
            "ASGS_2021": ["501011001", "501011001", "5", "9"],
            "TIME_PERIOD": ["2024", "2023", "2024", "2024"],
            "OBS_VALUE": ["12", "300", "250000", "50"],
        }
    )
    t = schema.make(abs_regional.tidy(raw, "ASGS_2021", "test"), source="abs_regional", adjustment="original")
    assert len(t) == 3  # Other Territories dropped
    assert tuple(t.iloc[0][["measure", "industry_code", "geo_level"]]) == ("businesses", "G", "SA2")
    assert t.iloc[0]["period"] == pd.Timestamp("2024-06-30") and t.iloc[0]["period_type"] == "day"
    assert t.iloc[1]["period"] == pd.Timestamp("2022-07-01") and t.iloc[1]["period_type"] == "fy"
    assert tuple(t.iloc[2][["geo_level", "geo_code", "industry_level"]]) == ("STE", "WA", "total")


def test_state_of_new_levels():
    assert schema.state_of("IVI_REGION", "NSW:Hunter") == "NSW"
    assert schema.state_of("CLUE_AREA", "Carlton") == "VIC"
    assert schema.state_of("LGA", "24600") == "VIC"


def _time_series_workbook(path, series: dict[str, tuple[str, list[float]]], dates: list[dt.datetime]) -> None:
    """A minimal ABS time series workbook: Index sheet and one Data sheet with the ten header rows."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Index"
    ws = wb.create_sheet("Data1")
    ids = list(series)
    header = [
        [None, *[series[i][0] for i in ids]],
        ["Unit", *["$ Millions"] * len(ids)],
        ["Series Type", *["Original"] * len(ids)],
        ["Data Type", *["DERIVED"] * len(ids)],
        ["Frequency", *["Annual"] * len(ids)],
        ["Collection Month", *[6] * len(ids)],
        ["Series Start", *[dates[0]] * len(ids)],
        ["Series End", *[dates[-1]] * len(ids)],
        ["No. Obs", *[len(dates)] * len(ids)],
        ["Series ID", *ids],
    ]
    for row in header:
        ws.append(row)
    for k, d in enumerate(dates):
        ws.append([d, *[series[i][1][k] for i in ids]])
    wb.save(path)


def test_state_accounts_reader_keeps_divisions_and_totals(tmp_path):
    from alicia.pipeline import abs_state_accounts, abs_tsw

    path = tmp_path / "wa.xlsx"
    dates = [dt.datetime(2023, 6, 1), dt.datetime(2024, 6, 1)]
    _time_series_workbook(
        path,
        {
            "S1": ("Mining (B) ;  Industry gross value added: Chain volume measures ;", [100.0, 110.0]),
            "S2": ("Mining (B) ;  Industry gross value added: Current prices ;", [90.0, 120.0]),
            "S3": ("Ownership of dwellings ;  Industry gross value added: Current prices ;", [5.0, 6.0]),
            "S4": ("Total all industries ;  Industry gross value added: Current prices ;", [200.0, 230.0]),
            "S5": ("Mining (B) ;  Industry gross value added: Chain volume measures - Percentage changes ;", [1, 10]),
        },
        dates,
    )
    raw = abs_tsw.read(path)
    assert set(raw["series_id"]) == {"S1", "S2", "S3", "S4", "S5"} and len(raw) == 10
    out = abs_state_accounts.tidy_industry(raw, "WA", "test")
    panel = schema.make(out, source="abs_state_accounts", industry_scheme="ANZSIC06")
    assert sorted(panel["measure"].unique()) == ["gva_current", "gva_volume"]
    assert set(panel["industry_code"]) == {"B", ""}  # ownership of dwellings and percentage changes dropped
    fy = panel.loc[(panel["measure"] == "gva_current") & (panel["industry_code"] == "B"), "period"]
    assert list(fy.dt.strftime("%Y-%m-%d")) == ["2022-07-01", "2023-07-01"]  # year ending June -> FY start


def test_rba_i2_reader(tmp_path):
    import openpyxl

    from alicia.pipeline import controls

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["I2 COMMODITY PRICES", None, None])
    ws.append(["Title", "Commodity prices – A$", "Rural commodity prices – A$"])
    ws.append(["Description", "x", "y"])
    ws.append(["Frequency", "Monthly", "Monthly"])
    ws.append(["Type", "Original", "Original"])
    ws.append(["Units", "Index, 2024/25=100", "Index, 2024/25=100"])
    ws.append([None, None, None])
    ws.append([None, None, None])
    ws.append(["Source", "RBA", "RBA"])
    ws.append(["Publication date", dt.datetime(2026, 10, 1), dt.datetime(2026, 10, 1)])
    ws.append(["Series ID", "GRCPAIAD", "GRCPRCAD"])
    ws.append([dt.datetime(2026, 7, 31), 101.5, 99.0])
    ws.append([dt.datetime(2026, 8, 31), 102.0, None])
    path = tmp_path / "i2.xlsx"
    wb.save(path)
    d = controls.read_rba_i2(path)
    assert len(d) == 3 and set(d["series_id"]) == {"GRCPAIAD", "GRCPRCAD"}
    assert d.loc[d["series_id"] == "GRCPAIAD", "value"].tolist() == [101.5, 102.0]


def test_controls_are_kept_out_of_the_outcome_catalogue(tmp_path, monkeypatch):
    from alicia.pipeline import store

    monkeypatch.setattr(store, "PROCESSED", tmp_path)
    row = schema.make(
        pd.DataFrame(
            {
                "measure": ["population"],
                "industry_level": ["total"],
                "geo_level": ["STE"],
                "geo_code": ["WA"],
                "period_type": ["quarter"],
                "period": ["2020-01-01"],
                "value": [1.0],
            }
        ),
        source="abs_erp",
        industry_scheme="none",
    )
    store.write_panel(row, "controls", "state_quarter")
    store.write_panel(row.assign(source="abs_x"), "abs_x", "outcome")
    assert [(s, t) for s, t, _ in store.panels()] == [("abs_x", "outcome")]
