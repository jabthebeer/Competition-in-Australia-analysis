"""Reader for ABS time series workbooks (the standard "Time Series Workbook" spreadsheet layout).

Each workbook has an Index sheet and one or more Data sheets. A Data sheet has one column per series
with ten header rows (description, unit, series type, data type, frequency, collection month,
series start, series end, number of observations, series ID) and then one row per period, dated at
the period's last month.
"""

from __future__ import annotations

import openpyxl
import pandas as pd

HEADER = ("description", "unit", "series_type", "data_type", "frequency", "collection_month", "start", "end", "obs")


def read(path) -> pd.DataFrame:
    """Long frame: series_id, description, unit, series_type, frequency, date, value."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    frames = []
    for name in wb.sheetnames:
        if not name.startswith("Data"):
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        if len(rows) < 11 or rows[9][0] != "Series ID":
            raise ValueError(f"{path}: sheet {name} is not in the time series workbook layout")
        meta = pd.DataFrame(
            {key: [r for r in rows[i][1:]] for i, key in enumerate(HEADER)} | {"series_id": list(rows[9][1:])}
        ).dropna(subset=["series_id"])
        body = pd.DataFrame([r[: len(meta) + 1] for r in rows[10:] if r[0] is not None])
        body.columns = ["date", *meta["series_id"]]
        long = body.melt(id_vars="date", var_name="series_id", value_name="value").dropna(subset=["value"])
        frames.append(long.merge(meta[["series_id", "description", "unit", "series_type", "frequency"]]))
    wb.close()
    out = pd.concat(frames, ignore_index=True)
    out["date"] = pd.to_datetime(out["date"])
    out["value"] = pd.to_numeric(out["value"], errors="coerce")
    return out.dropna(subset=["value"])
