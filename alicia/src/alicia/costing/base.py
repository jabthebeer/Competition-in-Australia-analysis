"""Shared parts of the costing parameter connectors: releases, downloads, tables and deterministic JSON.

A connector is a module (or any object) with:

- ``SOURCE``: a :class:`Source` (id, title, publisher, licence, landing page);
- ``discover() -> Release``: the latest release, read from the publisher's latest-release page or API;
- ``fetch(release)``: download the release's files to ``data/raw/<source>/`` and record them in
  ``data/manifests/<source>.json`` (URL, retrieval time, size, SHA-256);
- ``build(release) -> dict[str, Table]``: the tidy tables, read from those raw files only.

The per-source output (``data/processed/costing/sources/<source>.json``) holds the release, every raw
file's URL and SHA-256, and the tables. It carries no build timestamp: the only date that depends on when it
was made is each file's retrieval date, taken from the manifest. Re-running ``build`` on the same raw files
gives a byte-identical file.
"""

from __future__ import annotations

import datetime as dt
import json
import math
import re
import time
from dataclasses import dataclass, field
from html import unescape
from pathlib import Path
from urllib.parse import urljoin, urlsplit

from ..pipeline import store

SOURCE_SCHEMA = "alicia-costing-source"
PARAMS_SCHEMA = "alicia-costing-parameters"
VERSION = 1
ABS_LICENCE = "CC BY 4.0, Australian Bureau of Statistics"
ABS = "https://www.abs.gov.au"
SIGNIFICANT = 7  # significant digits kept for floats


def costing_dir() -> Path:
    return store.PROCESSED / "costing"


def sources_dir() -> Path:
    return costing_dir() / "sources"


@dataclass(frozen=True)
class Source:
    id: str
    title: str
    publisher: str
    licence: str
    url: str  # landing page (latest release, or the API's dataflow)


@dataclass(frozen=True)
class RemoteFile:
    name: str  # path under data/raw/<source>/, e.g. "2024-25/5204005_GVA_by_Industry.xlsx"
    url: str
    accept: str | None = None  # HTTP Accept header for API queries


@dataclass(frozen=True)
class Release:
    """One release of a source, as the publisher states it."""

    id: str  # the publisher's release identifier (reference period, version or API vintage)
    date: str | None  # ISO release date, if the publisher states one
    files: tuple[RemoteFile, ...]
    page: str  # where the release was discovered (latest-release page or API)
    next_release: str | None = None  # ISO date, if stated
    extra: tuple[tuple[str, str], ...] = ()  # facts read at discovery, e.g. the chain volume reference year

    def get(self, key: str, default: str | None = None) -> str | None:
        return dict(self.extra).get(key, default)

    def fingerprint(self) -> dict:
        return {"id": self.id, "date": self.date, "files": sorted(f.name for f in self.files)}

    def to_json(self) -> dict:
        return {
            "id": self.id,
            "date": self.date,
            "next_release": self.next_release,
            "page": self.page,
            "files": [{"name": f.name, "url": f.url} | ({"accept": f.accept} if f.accept else {}) for f in self.files],
            "extra": dict(self.extra),
        }

    @classmethod
    def from_json(cls, d: dict) -> Release:
        return cls(
            id=d["id"],
            date=d.get("date"),
            files=tuple(RemoteFile(f["name"], f["url"], f.get("accept")) for f in d["files"]),
            page=d["page"],
            next_release=d.get("next_release"),
            extra=tuple(sorted(d.get("extra", {}).items())),
        )


def is_newer(release: Release, committed: dict | None) -> bool:
    """True when the discovered release differs from the committed one and is not older than it."""
    if committed is None:
        return True
    old = Release.from_json(committed["release"])
    if release.fingerprint() == old.fingerprint():
        return False
    return not (release.date and old.date and release.date < old.date)


# ---------------------------------------------------------------- tables


@dataclass
class Table:
    """A tidy table: every row has exactly the declared columns; ``keys`` identify a row (sorted, unique)."""

    description: str
    columns: dict[str, str]  # column -> unit (or "text", "code: ..." for identifiers)
    keys: tuple[str, ...]
    rows: list[dict] = field(default_factory=list)

    def finish(self) -> Table:
        """Round floats, sort rows by their keys and check the shape. Returns self."""
        names = set(self.columns)
        missing = [k for k in self.keys if k not in names]
        if missing:
            raise ValueError(f"keys {missing} are not columns")
        if not self.rows:
            raise ValueError(f"table '{self.description[:40]}' has no rows")
        clean = []
        for r in self.rows:
            if set(r) != names:
                raise ValueError(f"row columns {sorted(r)} differ from declared {sorted(names)}")
            clean.append({k: num(v) for k, v in r.items()})
        clean.sort(key=lambda r: tuple(_sort_key(r[k]) for k in self.keys))
        seen = set()
        for r in clean:
            key = tuple(r[k] for k in self.keys)
            if key in seen:
                raise ValueError(f"duplicate key {key} in table '{self.description[:40]}'")
            seen.add(key)
        self.rows = clean
        return self

    def to_json(self) -> dict:
        return {"description": self.description, "columns": dict(self.columns), "rows": self.rows}


def _sort_key(v):
    return (0, "") if v is None else (1, v) if isinstance(v, str) else (2, v)


def num(v):
    """Floats rounded to SIGNIFICANT digits (integral values become ints); NaN becomes None."""
    if isinstance(v, bool) or v is None or isinstance(v, (str, int)):
        return v
    try:
        x = float(v)
    except (TypeError, ValueError):
        return v
    if math.isnan(x) or math.isinf(x):
        return None
    if x == 0:
        return 0
    r = float(f"{x:.{SIGNIFICANT}g}")
    return int(r) if r.is_integer() and abs(r) < 1e15 else r


def value(cell) -> float | None:
    """A numeric workbook cell; markers such as 'na', 'np', '..', '-' and blanks become None."""
    if cell is None or isinstance(cell, bool):
        return None
    if isinstance(cell, (int, float)):
        return None if isinstance(cell, float) and math.isnan(cell) else float(cell)
    s = str(cell).strip().replace(",", "")
    s = re.sub(r"[\^*#]+$", "", s)  # RSE and footnote markers appended to a value
    try:
        return float(s)
    except ValueError:
        return None


def clean_text(s) -> str:
    return re.sub(r"\s+", " ", str(s).replace("\xa0", " ")).strip() if s is not None else ""


# ---------------------------------------------------------------- ANZSIC names


def _norm(s: str) -> str:
    s = clean_text(s).lower().replace("&", " and ")
    s = re.sub(r"\(private\)|\(incl\.?[^)]*\)", " ", s)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return " ".join(w for w in s.split() if w not in {"and", "the"})


def anzsic_codes() -> dict[str, dict]:
    from .. import anzsic

    return anzsic._codes()


def division_by_name(name: str) -> str | None:
    """Division letter for a division title in any capitalisation ('Information, media and telecommunications')."""
    target = _norm(name)
    for code, c in anzsic_codes().items():
        if len(code) == 1 and _norm(c["name"]) == target:
            return code
    return None


def child_by_name(parent: str, name: str) -> str | None:
    """Code of the child of ``parent`` (a division or subdivision) whose title matches ``name``."""
    target = _norm(name)
    for code, c in anzsic_codes().items():
        if c["parent"] == parent and _norm(c["name"]) == target:
            return code
    return None


# ---------------------------------------------------------------- HTTP and discovery


_CLIENT = None


def client():
    """One shared HTTP client (descriptive User-Agent, redirects followed)."""
    global _CLIENT
    import httpx

    if _CLIENT is None:
        _CLIENT = httpx.Client(headers={"User-Agent": store.USER_AGENT}, timeout=180, follow_redirects=True)
    return _CLIENT


def get(url: str, accept: str | None = None, tries: int = 3):
    """GET with a few retries on transport errors and 5xx responses; raises on failure."""
    import httpx

    last = None
    for attempt in range(tries):
        try:
            resp = client().get(url, headers={"Accept": accept} if accept else None)
            if resp.status_code < 500:
                resp.raise_for_status()
                return resp
            last = httpx.HTTPStatusError(f"HTTP {resp.status_code} for {url}", request=resp.request, response=resp)
        except httpx.TransportError as exc:
            last = exc
        time.sleep(2 * (attempt + 1))
    raise last


def page_text(html: str) -> str:
    return re.sub(r"\s+", " ", unescape(re.sub(r"<[^>]+>", " ", html))).strip()


def links(html: str, base: str) -> list[str]:
    """Every href on a page, made absolute, in page order without repeats."""
    out = []
    for h in re.findall(r'href="([^"]+)"', html):
        u = urljoin(base, unescape(h))
        if u not in out:
            out.append(u)
    return out


def _dmy(s: str) -> str:
    d, m, y = (int(x) for x in re.split(r"\s*/\s*", s.strip()))
    return dt.date(y, m, d).isoformat()


@dataclass(frozen=True)
class AbsPage:
    url: str
    reference_period: str
    released: str  # ISO; the later of "Released" and any "Release date and time" update
    next_release: str | None
    links: tuple[str, ...]
    text: str

    def files(self, pattern: str) -> list[str]:
        rx = re.compile(pattern)
        return [u for u in self.links if rx.search(urlsplit(u).path)]


def abs_page(url: str) -> AbsPage:
    """Read an ABS release page: reference period, release date (and update), next release, links."""
    html = get(url).text
    text = page_text(html)
    m = re.search(r"Reference period (.+?) Released (\d{1,2}/\d{1,2}/\d{4})", text)
    if not m:
        raise ValueError(f"{url}: no 'Reference period ... Released dd/mm/yyyy' on the page")
    released = _dmy(m.group(2))
    upd = [_dmy(x) for x in re.findall(r"Release date and time (\d{1,2}/\d{1,2}/\d{4})", text)]
    nxt = re.search(r"Next [Rr]elease (\d{1,2}/\d{1,2}/\d{4})", text)
    return AbsPage(
        url=url,
        reference_period=m.group(1).strip(),
        released=max([released, *upd]),
        next_release=_dmy(nxt.group(1)) if nxt else None,
        links=tuple(links(html, ABS)),
        text=text,
    )


def release_slug(url: str) -> str:
    """The path segment before the file name: '.../2024-25/5204005_GVA_by_Industry.xlsx' -> '2024-25'."""
    parts = urlsplit(url).path.rstrip("/").split("/")
    return parts[-2]


def file_name(url: str) -> str:
    from urllib.parse import unquote

    return unquote(urlsplit(url).path.rsplit("/", 1)[-1])


def remote(url: str, folder: str | None = None, name: str | None = None, accept: str | None = None) -> RemoteFile:
    """A file kept at data/raw/<source>/<folder>/<name>; folder defaults to the URL's release slug."""
    folder = folder or release_slug(url)
    name = name or file_name(url)
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", name)
    return RemoteFile(f"{re.sub(r'[^A-Za-z0-9._-]+', '_', folder)}/{safe}", url, accept)


# ---------------------------------------------------------------- raw files and manifests


def local_path(source: str, f: RemoteFile) -> Path:
    return store.RAW / source / f.name


def fetch_files(source: Source, release: Release) -> list[Path]:
    """Download each file of a release once (kept under its release folder) and record it in the manifest."""
    import httpx

    paths = []
    for f in release.files:
        path = local_path(source.id, f)
        if not path.exists():
            path.parent.mkdir(parents=True, exist_ok=True)
            tmp = path.with_name(path.name + ".part")
            headers = {"User-Agent": store.USER_AGENT} | ({"Accept": f.accept} if f.accept else {})
            for attempt in range(3):
                try:
                    with (
                        httpx.Client(headers=headers, timeout=300, follow_redirects=True) as c,
                        c.stream("GET", f.url) as resp,
                    ):
                        resp.raise_for_status()
                        with tmp.open("wb") as out:
                            for chunk in resp.iter_bytes(1 << 20):
                                out.write(chunk)
                    break
                except httpx.TransportError:
                    if attempt == 2:
                        raise
                    time.sleep(3 * (attempt + 1))
            tmp.replace(path)
            store.record_download(source.id, path, f.url, source.licence, note=f"release {release.id}")
        paths.append(path)
    return paths


def manifest_entry(source: str, path: Path) -> dict | None:
    mpath = store.MANIFESTS / f"{source}.json"
    if not mpath.exists():
        return None
    rel = str(path.relative_to(store.ROOT))
    for e in json.loads(mpath.read_text("utf-8")).get("files", []):
        if e["file"] == rel:
            return e
    return None


def file_records(source: Source, release: Release) -> list[dict]:
    """URL, SHA-256, size and retrieval date of each raw file (the hash is computed from the file itself)."""
    out = []
    for f in release.files:
        path = local_path(source.id, f)
        if not path.exists():
            raise FileNotFoundError(f"{path.relative_to(store.ROOT)} is missing; run 'alicia costing-params --refresh'")
        entry = manifest_entry(source.id, path) or {}
        sha = store.sha256(path)
        retrieved = entry.get("retrieved", "")[:10] if entry.get("sha256") == sha else None
        out.append({"name": f.name, "url": f.url, "sha256": sha, "bytes": path.stat().st_size, "retrieved": retrieved})
    return out


def have_files(source: Source, release: Release) -> bool:
    return all(local_path(source.id, f).exists() for f in release.files)


# ---------------------------------------------------------------- deterministic JSON


def _j(v) -> str:
    return json.dumps(v, sort_keys=True, ensure_ascii=False, separators=(",", ":"), allow_nan=False)


LINE_LISTS = {"rows", "sources", "records"}  # lists written one element per line, for readable diffs
LINE_DICTS = {None, "tables"}  # objects written one member per line


def dumps(obj, _key=None) -> str:
    """Compact JSON with sorted keys, one table row (or source, or record) per line. Deterministic."""
    if isinstance(obj, dict):
        items = [f"{_j(k)}:{dumps(obj[k], k)}" for k in sorted(obj)]
        if _key in LINE_DICTS:
            return "{\n" + ",\n".join(items) + "\n}"
        return "{" + ",".join(items) + "}"
    if isinstance(obj, list) and _key in LINE_LISTS:
        return "[\n" + ",\n".join(_j(x) for x in obj) + "\n]" if obj else "[]"
    return _j(obj)


def write_json(path: Path, obj) -> Path:
    text = dumps(obj) + "\n"
    if json.loads(text) != obj:  # pragma: no cover - guards the hand-written serialiser
        raise AssertionError("serialised JSON does not round-trip")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def source_doc(
    source: Source,
    release: Release,
    tables: dict[str, Table],
    derived_from: list | None = None,
    retrieved: str | None = None,
) -> dict:
    """The per-source document: release, raw files and tables. ``retrieved`` stands in for the files' retrieval
    date when the source has no raw files of its own (a bundled codelist, a processed panel)."""
    files = file_records(source, release)
    dates = [f["retrieved"] for f in files if f["retrieved"]] + ([retrieved] if retrieved else [])
    doc = {
        "schema": SOURCE_SCHEMA,
        "version": VERSION,
        "source": {
            "id": source.id,
            "title": source.title,
            "publisher": source.publisher,
            "licence": source.licence,
            "url": source.url,
        },
        "release": release.to_json(),
        "files": files,
        "retrieved": max(dates) if dates else None,
        "tables": {name: t.finish().to_json() for name, t in sorted(tables.items())},
    }
    if derived_from:
        doc["derived_from"] = derived_from
    return doc


def source_path(source_id: str) -> Path:
    return sources_dir() / f"{source_id}.json"


def read_source_doc(source_id: str) -> dict | None:
    p = source_path(source_id)
    return json.loads(p.read_text("utf-8")) if p.exists() else None


def sheet_rows(path: Path, sheet: str) -> list[list]:
    """Every row of a worksheet as a list of cell values (read only, formulas as last calculated)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()


def sheet_names(path: Path) -> list[str]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def fy_label(year_end: int) -> str:
    """2025 -> '2024-25'."""
    return f"{year_end - 1}-{year_end % 100:02d}"


def fy_end(label: str) -> int:
    """'2024-25' -> 2025."""
    m = re.fullmatch(r"(\d{4})\s*[-–]\s*(\d{2})", clean_text(label))
    if not m:
        raise ValueError(f"not a financial year: {label!r}")
    return int(m.group(1)) + 1
